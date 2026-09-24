import csv
import html
import io
from pathlib import Path
from .database import DATA_DIR

FORMATS = {
    "CSV": ("csv", "text/csv"), "Excel": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "HTML dashboard": ("html", "text/html"), "Report": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "Slides": ("pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"), "PDF": ("pdf", "application/pdf"),
}

def generate(project, points, format, export_id):
    ext, mime = FORMATS[format]
    directory = DATA_DIR / "exports"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{export_id}.{ext}"
    headers = ["Code", "Data point", "Owner", "Status", "Value", "Due", "Source / audit trail"]
    rows = [[p["code"], p["name"], p["owner"], p["status"], p.get("value") or "Not collected", p["due"], "; ".join(f'{h["action"]}: {h.get("reason", "")} ({h["at"]})' for h in p.get("history", []))] for p in points]
    # Spreadsheet apps interpret formula prefixes even in quoted CSV cells.
    safe_rows = [["'" + str(v) if str(v).startswith(("=", "+", "-", "@")) else str(v) for v in row] for row in rows]
    if format == "CSV":
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f); writer.writerow(headers); writer.writerows(safe_rows)
    elif format == "Excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "Data points"
        ws.append(headers)
        for row in safe_rows: ws.append(row)
        for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="2F4BFF")
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for col in "ABCDEFG": ws.column_dimensions[col].width = 30 if col != "G" else 65
        wb.save(path)
    elif format == "HTML dashboard":
        table = "".join("<tr>" + "".join(f"<td>{html.escape(str(v))}</td>" for v in row) + "</tr>" for row in rows)
        path.write_text(f'<!doctype html><html lang="en"><meta charset="utf-8"><title>{html.escape(project["name"])}</title><style>body{{font:14px system-ui;background:#F7F7F5;margin:40px;color:#1F1F1E}}table{{border-collapse:collapse;width:100%;background:white}}td,th{{border:1px solid #E6E6E3;padding:12px;text-align:left}}th{{background:#EBEEFF}}</style><h1>{html.escape(project["name"])}</h1><p>Oneput · Read-only data snapshot. Missing values remain explicitly marked.</p><table><tr>{"".join(f"<th>{h}</th>" for h in headers)}</tr>{table}</table></html>', encoding="utf-8")
    elif format == "Report":
        from docx import Document
        doc = Document(); doc.add_heading(project["name"], 0); doc.add_paragraph("Oneput · Draft report. Missing data is explicitly marked.")
        for row in rows:
            doc.add_heading(f"{row[0]} · {row[1]}", 2)
            doc.add_paragraph(f"Owner: {row[2]} | Status: {row[3]} | Due: {row[5]}")
            doc.add_paragraph(row[4]); doc.add_paragraph(row[6] or "No review history yet.")
        doc.save(path)
    elif format == "Slides":
        from pptx import Presentation
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text = project["name"]; slide.placeholders[1].text = "Oneput · Data collection snapshot"
        for row in rows:
            slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = f"{row[0]} · {row[1]}"
            slide.placeholders[1].text = f"Owner: {row[2]}\nStatus: {row[3]}\nValue: {row[4]}\nDue: {row[5]}"
        prs.save(path)
    else:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        styles = getSampleStyleSheet(); story = [Paragraph(html.escape(project["name"]), styles["Title"])]
        for row in rows:
            story += [Paragraph(html.escape(f"{row[0]} · {row[1]}"), styles["Heading2"]), Paragraph(html.escape(f"{row[3]} · {row[4]}"), styles["BodyText"]), Spacer(1, 12)]
        SimpleDocTemplate(str(path)).build(story)
    return path, mime
